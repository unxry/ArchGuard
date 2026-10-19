"""Build the editable thesis from public manuscript and verified bibliographic sources."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "docs/thesis"
TITLE = (
    "Разработка гибридного метода автоматизированного контроля соответствия "
    "программной архитектуры исходному коду на основе статического, графового "
    "и интеллектуального анализа"
)
AUTHOR = "Бендас Роман Александрович"
DEPARTMENT = "кафедра программной инженерии"
GROUP = "ПО-51м"
SUPERVISOR = "Петрик Е. А."
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
CITATION = re.compile(r"\[@([\w-]+(?:;\s*@[\w-]+)*)\]")
MARKER = re.compile(r"<!--\s*(FIGURE|TABLE):([\w-]+)\s*\|\s*(.*?)\s*-->")
REFERENCE = re.compile(r"\[(рисунок|таблица):([\w-]+)\]")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_files(manifest: dict[str, Any]) -> list[Path]:
    result = [THESIS / "source" / name for name in manifest["manuscript"]]
    if any(not p.is_file() for p in result):
        raise ValueError("Incomplete thesis manuscript")
    return result


def citation_order(text: str, library: dict[str, dict[str, Any]]) -> dict[str, int]:
    order: dict[str, int] = {}
    for match in CITATION.finditer(text):
        for key in re.findall(r"[\w-]+", match[1]):
            if key not in library or library[key].get("verified") is not True:
                raise ValueError("Unverified or missing bibliography source: " + key)
            order.setdefault(key, len(order) + 1)
    return order


def publication_guard(text: str) -> None:
    forbidden = r"sc-[a-f0-9]{24}|/Us" + r"ers/|sk-[A-Za-z0-9]{16,}|Unit[a-f0-9]{8}"
    if re.search(forbidden, text):
        raise ValueError("Private data or absolute host path in thesis")
    if "UI_SCREENSHOTS_PENDING" in text:
        raise ValueError("Internal UI placeholder cannot enter Word")


class ThesisBuilder:
    def __init__(self, manifest: dict[str, Any], pages: int | None = None) -> None:
        from docx import Document
        from docx.shared import Cm, Pt

        self.document: Any = Document()
        self.manifest = manifest
        self.Cm, self.Pt = Cm, Pt
        self.pagination = (
            read_json(THESIS / "internal/pagination.json")
            if (THESIS / "internal/pagination.json").exists()
            else {}
        )
        self.pages = pages or self.pagination.get("pages", 0)
        self.files = source_files(manifest)
        self.text = "\n\n".join(p.read_text(encoding="utf-8") for p in self.files)
        publication_guard(self.text)
        self.library = {r["key"]: r for r in read_json(THESIS / "internal/references.json")}
        self.citations = citation_order(self.text, self.library)
        self.figures: dict[str, dict[str, Any]] = {}
        self.tables: dict[str, dict[str, Any]] = {}
        self.headings: list[dict[str, Any]] = []
        self.equations = 0
        self.last_heading = ""
        self.pending_table: dict[str, Any] | None = None
        self.appendix = ""
        self._collect()
        self._styles()
        self._section(self.document.sections[0], first=True)

    @staticmethod
    def element(tag: str, **attributes: str) -> Any:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        node = OxmlElement(tag)
        for key, value in attributes.items():
            node.set(qn(key), value)
        return node

    def _collect(self) -> None:
        appendix = ""
        counters: dict[str, int] = {}
        for path in self.files:
            for line in path.read_text(encoding="utf-8").splitlines():
                if line.startswith("# "):
                    title = line[2:]
                    match = re.match(r"ПРИЛОЖЕНИЕ ([А-Я])", title)
                    appendix = match[1] if match else ""
                if re.match(r"#{1,3} ", line):
                    level = len(line.split(" ")[0])
                    title = line[level + 1 :]
                    self.headings.append(
                        dict(
                            title=title,
                            level=level,
                            bookmark=f"heading_{len(self.headings) + 1}",
                            appendix=appendix,
                        )
                    )
                marker = MARKER.fullmatch(line.strip())
                if marker:
                    kind, key, caption = marker.groups()
                    counter_key = kind + appendix
                    counters[counter_key] = counters.get(counter_key, 0) + 1
                    number = (appendix + "." if appendix else "") + str(counters[counter_key])
                    target = self.figures if kind == "FIGURE" else self.tables
                    if key in target:
                        raise ValueError("Duplicate figure/table key: " + key)
                    target[key] = dict(
                        key=key,
                        caption=caption,
                        number=number,
                        bookmark=("fig_" if kind == "FIGURE" else "tbl_") + key.replace("-", "_"),
                        sequence=("Figure" if kind == "FIGURE" else "Table") + appendix,
                        ordinal=counters[counter_key],
                        appendix=appendix,
                    )

    def _styles(self) -> None:
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import RGBColor

        styles = self.document.styles
        settings = {
            "Normal": (14, False, 1.5, 1.5, 0, 0),
            "Title": (14, False, 1.0, 0, 0, 0),
            "TitlePage": (14, False, 1.0, 0, 0, 5),
            "Assignment": (14, False, 1.0, 0, 0, 4),
            "Structural": (14, True, 1.5, 0, 0, 22.68),
            "FrontHeading": (14, True, 1.5, 0, 0, 22.68),
            "Heading 1": (14, True, 1.5, 1.5, 0, 22.68),
            "Heading 2": (14, True, 1.5, 1.5, 22.68, 22.68),
            "Heading 3": (14, True, 1.5, 1.5, 22.68, 22.68),
            "Caption": (14, False, 1.0, 0, 6, 6),
            "TableText": (12, False, 1.0, 0, 0, 0),
            "Equation": (14, False, 1.5, 0, 6, 6),
            "Bibliography": (14, False, 1.5, 1.5, 0, 0),
            "AppendixHeading": (14, True, 1.5, 0, 0, 22.68),
            "CodeExample": (10, False, 1.0, 0, 3, 3),
            "TOC 1": (14, False, 1.5, 0, 0, 0),
            "TOC 2": (14, False, 1.5, 0, 0, 0),
            "TOC 3": (14, False, 1.5, 0, 0, 0),
        }
        for name, (size, bold, spacing, indent, before, after) in settings.items():
            style = (
                styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            )
            style.base_style = styles["Normal"] if name != "Normal" else None
            style.font.name = "Times New Roman"
            style.font.size = self.Pt(size)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
            rpr = style.element.get_or_add_rPr()
            fonts = self.element(
                "w:rFonts",
                **{
                    "w:ascii": "Times New Roman",
                    "w:hAnsi": "Times New Roman",
                    "w:cs": "Times New Roman",
                    "w:eastAsia": "Times New Roman",
                },
            )
            rpr.append(fonts)
            rpr.append(self.element("w:lang", **{"w:val": "ru-RU"}))
            paragraph = style.paragraph_format
            paragraph.line_spacing = spacing
            paragraph.first_line_indent = self.Cm(indent)
            paragraph.space_before = self.Pt(before)
            paragraph.space_after = self.Pt(after)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.widow_control = True
            if name in {"Structural", "FrontHeading", "Title", "TitlePage", "AppendixHeading"}:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if name.startswith("Heading") or name in {
                "Structural",
                "AppendixHeading",
                "FrontHeading",
            }:
                paragraph.keep_with_next = True
            if name in {"Heading 1", "Structural", "AppendixHeading"}:
                paragraph.page_break_before = True
            if name.startswith("TOC"):
                level = int(name[-1])
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.left_indent = self.Cm((level - 1) * 0.65)
                paragraph.right_indent = self.Cm(0.8)
                paragraph.tab_stops.add_tab_stop(self.Cm(16.5), 2)
            if name == "CodeExample":
                style.font.name = "Courier New"
                for fonts_node in rpr.findall("{" + W + "}rFonts"):
                    for name_attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
                        fonts_node.set("{" + W + "}" + name_attribute, "Courier New")
            if name in {"Assignment", "CodeExample"}:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        settings_node = self.document.settings.element
        for border in styles.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        settings_node.append(self.element("w:updateFields", **{"w:val": "true"}))
        properties = self.document.core_properties
        properties.title = TITLE
        properties.subject = "Магистерская выпускная квалификационная работа"
        properties.author = ""
        properties.last_modified_by = ""
        properties.comments = ""
        properties.keywords = "ArchGuard AI, архитектура, гибридный анализ"
        properties.created = datetime(2026, 10, 16, tzinfo=UTC)
        properties.modified = properties.created

    def _section(self, section: Any, first: bool = False, landscape: bool = False) -> None:
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width = self.Cm(29.7 if landscape else 21)
        section.page_height = self.Cm(21 if landscape else 29.7)
        section.left_margin, section.right_margin = self.Cm(3), self.Cm(1.5)
        section.top_margin = section.bottom_margin = self.Cm(2)
        section.header_distance = self.Cm(0.8)
        section.footer_distance = self.Cm(1)
        section.different_first_page_header_footer = first
        if first:
            paragraph = section.footer.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = self.Cm(0)
            self.field(paragraph, "PAGE", "1")

    def field(self, paragraph: Any, instruction: str, result: str) -> None:
        run = paragraph.add_run()
        run._r.append(self.element("w:fldChar", **{"w:fldCharType": "begin"}))
        code = self.element("w:instrText")
        code.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        code.text = " " + instruction + " "
        run._r.append(code)
        run._r.append(self.element("w:fldChar", **{"w:fldCharType": "separate"}))
        paragraph.add_run(result)
        paragraph.add_run()._r.append(self.element("w:fldChar", **{"w:fldCharType": "end"}))

    def bookmark(self, paragraph: Any, name: str) -> None:
        index = str(len(self.document.element.xpath(".//w:bookmarkStart")) + 1)
        paragraph._p.insert(
            0, self.element("w:bookmarkStart", **{"w:id": index, "w:name": name[:40]})
        )
        paragraph._p.append(self.element("w:bookmarkEnd", **{"w:id": index}))

    def inline(self, paragraph: Any, text: str) -> None:
        text = re.sub(r"(?<=\d)\s*%", "\u00a0%", text)
        pattern = re.compile(
            r"\[@([\w-]+(?:;\s*@[\w-]+)*)\]|\[(рисунок|таблица):([\w-]+)\]|"
            r"\*\*(.+?)\*\*|`([^`]+)`"
        )
        start = 0
        for match in pattern.finditer(text):
            paragraph.add_run(text[start : match.start()])
            if match[1]:
                numbers = [str(self.citations[key]) for key in re.findall(r"[\w-]+", match[1])]
                paragraph.add_run("[" + "; ".join(numbers) + "]")
            elif match[2]:
                entry = (self.figures if match[2] == "рисунок" else self.tables)[match[3]]
                preceding = text[: match.start()]
                if not re.search(r"(?:рисунк|таблиц)\w*\s*$", preceding, re.IGNORECASE):
                    paragraph.add_run("рисунке " if match[2] == "рисунок" else "таблице ")
                self.field(paragraph, "REF " + entry["bookmark"] + " \\h", entry["number"])
            else:
                run = paragraph.add_run(match[4] or match[5])
                run.bold = bool(match[4])
            start = match.end()
        paragraph.add_run(text[start:])

    def paragraph(self, text: str, style: str = "Normal") -> Any:
        paragraph = self.document.add_paragraph(style=style)
        previous = paragraph._p.getprevious()
        if style == "Normal" and previous is not None and previous.tag == "{" + W + "}tbl":
            paragraph.paragraph_format.space_before = self.Pt(8)
        if style == "CodeExample":
            paragraph.add_run(text)
        else:
            self.inline(paragraph, text)
        return paragraph

    def title_page(self) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        for line in (
            "МИНОБРНАУКИ РОССИИ",
            "Федеральное государственное бюджетное образовательное учреждение\nвысшего образования",
            "«Юго-Западный государственный университет»",
            DEPARTMENT,
        ):
            self.paragraph(line, "TitlePage")
        p = self.paragraph("ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА\nПО ПРОГРАММЕ МАГИСТРАТУРЫ", "Title")
        p.paragraph_format.space_before = self.Pt(60)
        self.paragraph("[ЗАПОЛНИТЬ: КОД И НАПРАВЛЕНИЕ]\n[ЗАПОЛНИТЬ: ПРОФИЛЬ]", "TitlePage")
        self.paragraph("на тему", "TitlePage")
        p = self.paragraph("«" + TITLE + "»", "Title")
        p.paragraph_format.space_before = self.Pt(14)
        p.paragraph_format.space_after = self.Pt(12)
        self.paragraph("[ЗАПОЛНИТЬ: ВИД ВКР ПО ПРОГРАММЕ ГИА]", "TitlePage")
        for line in (
            "Автор ВКР: " + AUTHOR + "\nПодпись __________ Дата __________",
            "Группа: " + GROUP,
            "Руководитель ВКР: " + SUPERVISOR + "\nПодпись __________ Дата __________",
            "Нормоконтроль: [ЗАПОЛНИТЬ: НОРМОКОНТРОЛЬ]\nПодпись __________ Дата __________",
            "Рецензент: [ЗАПОЛНИТЬ: РЕЦЕНЗЕНТ]\nПодпись __________ Дата __________",
            "ВКР допущена к защите:",
            "Заведующий кафедрой: [ЗАПОЛНИТЬ: ЗАВЕДУЮЩИЙ КАФЕДРОЙ]"
            "\nПодпись __________ Дата __________",
        ):
            self.paragraph(line, "TitlePage").alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p = self.paragraph("Курск 2026", "TitlePage")
        p.paragraph_format.space_before = self.Pt(35)

    def assignment(self) -> None:
        self.document.add_page_break()
        for line in (
            "Юго-Западный государственный университет",
            DEPARTMENT,
            "УТВЕРЖДАЮ",
            "[ЗАПОЛНИТЬ: ЗАВЕДУЮЩИЙ КАФЕДРОЙ]",
            "[ЗАПОЛНИТЬ: ПОДПИСЬ И ДАТА УТВЕРЖДЕНИЯ]",
        ):
            self.paragraph(line, "TitlePage")
        self.paragraph(
            "ЗАДАНИЕ\nНА ВЫПУСКНУЮ КВАЛИФИКАЦИОННУЮ РАБОТУ\nПО ПРОГРАММЕ МАГИСТРАТУРЫ",
            "FrontHeading",
        )
        for line in (
            "Обучающийся: " + AUTHOR,
            "Шифр: [ЗАПОЛНИТЬ: ШИФР СТУДЕНТА]. Группа: " + GROUP,
            "1 Тема работы: «" + TITLE + "».",
            "Утверждена приказом ректора: [ЗАПОЛНИТЬ: ДАТА И НОМЕР ПРИКАЗА]",
            "2 Срок представления ВКР к защите: [ЗАПОЛНИТЬ: СРОК ПРЕДСТАВЛЕНИЯ К ЗАЩИТЕ]",
            "3 Исходные данные: исходный код ArchGuard AI; архитектурные правила; "
            "публичные протоколы, сводные результаты и реестры научных выводов; литература.",
            "4 Перечень вопросов, подлежащих исследованию (разработке): "
            "анализ методов; гибридный метод; "
            "статический, графовый и семантический анализ; формирование контекста; "
            "объединение свидетельств; прототип; экспериментальная оценка; отдельное "
            "расширение анализа безопасности; воспроизводимость и ограничения.",
            "5 Перечень графического материала: структура метода и прототипа; "
            "модель архитектуры и зависимостей; формирование контекста; механизм решения; "
            "схема эксперимента; сравнение методов и абляций; расширение безопасности.",
            "Руководитель: " + SUPERVISOR + "; подпись _________ дата _________",
            "Задание принял к исполнению: подпись студента _________ дата _________",
        ):
            self.paragraph(line, "Assignment")

    def abstract(self) -> None:
        self.document.add_page_break()
        self.paragraph("РЕФЕРАТ", "FrontHeading")
        p = self.paragraph("Объем работы — ")
        self.field(p, "NUMPAGES", str(self.pages) if self.pages else "0")
        p.add_run(
            f" с.; рисунков — {len(self.figures)}, таблиц — {len(self.tables)}, "
            f"{len([h for h in self.headings if h['title'].startswith('ПРИЛОЖЕНИЕ ')])} "
            f"приложений; источников — {len(self.citations)}."
        )
        self.paragraph(
            "архитектура программного обеспечения, статический анализ, граф зависимостей, "
            "языковая модель, гибридный метод, архитектурное соответствие, воспроизводимость"
        )
        self.paragraph(self.manifest["abstract"])

    def toc(self) -> None:
        self.document.add_page_break()
        self.paragraph("СОДЕРЖАНИЕ", "FrontHeading")
        p = self.document.add_paragraph()
        p.paragraph_format.space_after = self.Pt(0)
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.first_line_indent = self.Cm(0)
        run = p.add_run()
        run._r.append(self.element("w:fldChar", **{"w:fldCharType": "begin"}))
        code = self.element("w:instrText")
        code.text = ' TOC \\o "1-3" \\h \\z \\t "Structural,1,AppendixHeading,1" '
        run._r.append(code)
        run._r.append(self.element("w:fldChar", **{"w:fldCharType": "separate"}))
        entries = [h for h in self.headings if not h["appendix"]]
        entries += [
            dict(title="СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1, bookmark="bibliography")
        ]
        entries += [h for h in self.headings if h["appendix"]]
        for heading in entries:
            paragraph = self.paragraph(heading["title"], "TOC " + str(heading["level"]))
            paragraph.add_run("\u00a0\u00a0\t")
            page = str(self.pagination.get("headings", {}).get(heading["title"], "0"))
            self.field(paragraph, "PAGEREF " + heading["bookmark"] + " \\h", page)
        paragraph.add_run()._r.append(self.element("w:fldChar", **{"w:fldCharType": "end"}))

    def heading(self, text: str, level: int) -> None:
        from docx.enum.section import WD_SECTION_START

        self.last_heading = text
        entry = next(h for h in self.headings if h["title"] == text)
        is_appendix = text.startswith("ПРИЛОЖЕНИЕ ")
        self.appendix = entry["appendix"]
        landscape = is_appendix and self.appendix in self.manifest.get("landscape_appendices", [])
        new_section = is_appendix and (
            landscape
            or self.document.sections[-1].page_width > self.document.sections[-1].page_height
        )
        if new_section:
            section = self.document.add_section(WD_SECTION_START.NEW_PAGE)
            self._section(section, landscape=landscape)
        style = (
            "AppendixHeading"
            if is_appendix
            else "Structural"
            if level == 1 and not re.match(r"\d+ ", text)
            else "Heading " + str(level)
        )
        display = re.sub(r"^(ПРИЛОЖЕНИЕ [А-Я]) ", r"\1\n", text) if is_appendix else text
        paragraph = self.paragraph(display, style)
        if new_section:
            paragraph.paragraph_format.page_break_before = False
        self.bookmark(paragraph, entry["bookmark"])

    def caption(self, entry: dict[str, Any], kind: str) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        paragraph = self.document.add_paragraph(style="Caption")
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if kind == "Рисунок" else WD_ALIGN_PARAGRAPH.LEFT
        )
        paragraph.add_run(kind + " ")
        # Bookmark includes only the displayed number, not its explanatory caption.
        start = self.element(
            "w:bookmarkStart",
            **{
                "w:id": str(1000 + len(self.document.element.xpath(".//w:bookmarkStart"))),
                "w:name": entry["bookmark"][:40],
            },
        )
        end = self.element("w:bookmarkEnd", **{"w:id": start.get("{" + W + "}id")})
        paragraph._p.append(start)
        if entry["appendix"]:
            paragraph.add_run(entry["appendix"] + ".")
        self.field(paragraph, f"SEQ {entry['sequence']} \\* ARABIC", str(entry["ordinal"]))
        paragraph._p.append(end)
        paragraph.add_run(" — " + entry["caption"])
        paragraph.paragraph_format.keep_with_next = kind == "Таблица"

    def figure(self, key: str) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        entry = self.figures[key]
        assets = read_json(THESIS / "assets/figures.json")
        asset = assets[key]
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = self.Cm(0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(
            str(THESIS / "assets" / asset["png"]), width=self.Cm(asset.get("width_cm", 16.5))
        )
        drawing = paragraph._p.xpath(".//wp:docPr")[0]
        drawing.set("descr", entry["caption"])
        self.caption(entry, "Рисунок")

    def table(self, rows: list[list[str]], entry: dict[str, Any]) -> None:
        from PIL import ImageFont

        landscape = self.document.sections[-1].page_width > self.document.sections[-1].page_height
        width = 25.2 if landscape else 16.5
        weights: list[float] = [
            max(5, min(35, max(len(r[i]) for r in rows))) for i in range(len(rows[0]))
        ]
        if len(rows[0]) >= 6:
            weights = [max(4, min(20, len(c) + 3)) for c in rows[0]]
        if entry["key"].startswith("chronology_"):
            weights = [19, 40, 41]
        weights = self.manifest.get("table_column_weights", {}).get(entry["key"], weights)
        if any(re.fullmatch(r"(?:ARCH|SEC)\d{3}", row[0]) for row in rows[1:]):
            weights[0] = max(weights[0], 10)
        font: Any = ImageFont.truetype(
            self.manifest.get(
                "table_measurement_font", "/System/Library/Fonts/Supplemental/Times New Roman.ttf"
            ),
            120,
        )
        numeric = [
            all(re.fullmatch(r"[−-]?\d+(?:[,.]\d+)?|null", row[i]) for row in rows[1:])
            for i in range(len(rows[0]))
        ]
        if any(numeric):
            minimum = [
                max(float(font.getlength(row[i])) / 10 for row in rows[1:]) + 12
                if is_numeric
                else 42.0
                for i, is_numeric in enumerate(numeric)
            ]
            remaining = width * 28.3465 - sum(minimum)
            if remaining < 0:
                raise ValueError("Numeric table needs wider layout: " + entry["key"])
            total_weight = sum(weights)
            weights = [
                floor + remaining * weight / total_weight
                for floor, weight in zip(minimum, weights, strict=True)
            ]

        def height(row: list[str]) -> float:
            maximum = 1
            for value, weight in zip(row, weights, strict=True):
                capacity = width * 28.3465 * weight / sum(weights) - 9
                value = CITATION.sub("[99]", value).replace("<br>", "\n")
                count = 0
                for line in value.splitlines():
                    used, lines = 0.0, 1
                    for word in line.split():
                        size = float(font.getlength(word + " ")) / 10
                        if used and used + size > capacity:
                            lines += 1
                            used = 0
                        while size > capacity:
                            lines += 1
                            size -= capacity
                        used += size
                    count += lines
                maximum = max(maximum, count)
            return maximum * 14.5 + 12

        budget = 390 if landscape else 540
        first_budget = self.manifest.get("first_table_part_budget_pt", {}).get(entry["key"], budget)
        chunks: list[list[list[str]]] = []
        chunk, used = [rows[0]], height(rows[0])
        for row in rows[1:]:
            row_height = height(row)
            if used + row_height > (budget if chunks else first_budget) and len(chunk) > 1:
                chunks.append(chunk)
                chunk, used = [rows[0]], height(rows[0])
            chunk.append(row)
            used += row_height
        chunks.append(chunk)
        if entry["key"].startswith("chronology_") and len(chunks) == 3 and len(rows) == 7:
            balanced = [[rows[0], *rows[start : start + 2]] for start in (1, 3, 5)]
            if all(sum(height(row) for row in part) <= budget for part in balanced):
                chunks = balanced
        for index, chunk in enumerate(chunks):
            self._table_part(chunk, entry, weights, width, continuation=index > 0)

    def _table_part(
        self,
        rows: list[list[str]],
        entry: dict[str, Any],
        weights: list[float],
        width: float,
        continuation: bool,
    ) -> None:
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        if continuation:
            p = self.paragraph("Продолжение таблицы ", "Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.page_break_before = True
            self.field(p, "REF " + entry["bookmark"] + " \\h", entry["number"])
        else:
            self.caption(entry, "Таблица")
        columns = len(rows[0])
        table = self.document.add_table(rows=0, cols=columns)
        table.alignment = 1
        table.autofit = False
        for column, weight in zip(table.columns, weights, strict=True):
            column.width = self.Cm(width * weight / sum(weights))
        borders = self.element("w:tblBorders")
        for border in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(
                self.element("w:" + border, **{"w:val": "single", "w:sz": "4", "w:color": "000000"})
            )
        table._tbl.tblPr.append(borders)
        for index, source in enumerate(rows):
            row = table.add_row()
            row._tr.get_or_add_trPr().append(self.element("w:cantSplit"))
            if index == 0:
                row._tr.get_or_add_trPr().append(self.element("w:tblHeader"))
            for cell, value, weight in zip(row.cells, source, weights, strict=True):
                cell.width = self.Cm(width * weight / sum(weights))
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                paragraph = cell.paragraphs[0]
                paragraph.style = "TableText"
                paragraph.paragraph_format.widow_control = True
                paragraph.paragraph_format.keep_with_next = index < len(rows) - 1
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if index == 0 or re.fullmatch(
                    r"[−-]?\d+(?:[,.]\d+)?|[AА][1-7]|(?:ARCH|SEC)\d{3}|—", value
                ):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.inline(paragraph, value.replace("<br>", "\n"))
                if len(rows[0]) >= 6:
                    for run in paragraph.runs:
                        run.font.size = self.Pt(11)
                if index == 0:
                    cell._tc.get_or_add_tcPr().append(self.element("w:shd", **{"w:fill": "D9E2E8"}))
                    for run in paragraph.runs:
                        run.bold = True
                elif index % 2 == 0:
                    cell._tc.get_or_add_tcPr().append(self.element("w:shd", **{"w:fill": "F5F5F5"}))
                margins = self.element("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    margins.append(self.element("w:" + side, **{"w:w": "100", "w:type": "dxa"}))
                cell._tc.get_or_add_tcPr().append(margins)

    def equation(self, formula: str) -> None:
        self.equations += 1
        paragraph = self.document.add_paragraph(style="Equation")
        paragraph.paragraph_format.tab_stops.add_tab_stop(self.Cm(8.25), 1)
        paragraph.paragraph_format.tab_stops.add_tab_stop(self.Cm(16.5), 2)
        paragraph.add_run("\t")
        math = self.element("m:oMath")

        def text_node(text: str, upright: bool = False, bold: bool = False) -> Any:
            run = self.element("m:r")
            if upright or bold:
                properties = self.element("m:rPr")
                properties.append(self.element("m:sty", **{"m:val": "b" if bold else "p"}))
                run.append(properties)
            content = self.element("m:t")
            content.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            content.text = text
            run.append(content)
            return run

        tokens = re.findall(r"\\[A-Za-z]+|[{}_^]|[^\\{}_^]", formula)
        cursor = 0
        symbols = {"mu": "μ", "beta": "β", "lambda": "λ", "eta": "η", "sum": "∑"}

        def group() -> list[Any]:
            nonlocal cursor
            if cursor < len(tokens) and tokens[cursor] == "{":
                cursor += 1
                value = sequence(True)
                if cursor >= len(tokens) or tokens[cursor] != "}":
                    raise ValueError("Unbalanced equation group")
                cursor += 1
                return value
            return [atom()]

        def atom() -> Any:
            nonlocal cursor
            token = tokens[cursor]
            cursor += 1
            if not token.startswith("\\"):
                return text_node(token)
            command = token[1:]
            if command == "frac":
                fraction = self.element("m:f")
                for tag in ("m:num", "m:den"):
                    child = self.element(tag)
                    child.extend(group())
                    fraction.append(child)
                return fraction
            if command == "widetilde":
                accent = self.element("m:acc")
                properties = self.element("m:accPr")
                properties.append(self.element("m:chr", **{"m:val": "̃"}))
                value = self.element("m:e")
                value.extend(group())
                accent.extend([properties, value])
                return accent
            if command in {"text", "mathbf"}:
                value = group()
                content = "".join("".join(n.itertext()) for n in value)
                return text_node(content, upright=True, bold=command == "mathbf")
            if command in symbols:
                return text_node(symbols[command])
            if command in {"log", "exp"}:
                return text_node(command, upright=True)
            if command == "quad":
                return text_node(" ")
            raise ValueError("Unsupported equation command: " + command)

        def sequence(stop: bool = False) -> list[Any]:
            nonlocal cursor
            nodes: list[Any] = []
            while cursor < len(tokens) and (not stop or tokens[cursor] != "}"):
                if tokens[cursor] in {"_", "^"}:
                    operator = tokens[cursor]
                    cursor += 1
                    if not nodes:
                        raise ValueError("Equation script has no base")
                    base = nodes.pop()
                    script = self.element("m:sSub" if operator == "_" else "m:sSup")
                    child = self.element("m:e")
                    child.append(base)
                    value = self.element("m:sub" if operator == "_" else "m:sup")
                    value.extend(group())
                    script.extend([child, value])
                    nodes.append(script)
                elif tokens[cursor] == "{":
                    nodes.extend(group())
                else:
                    nodes.append(atom())
            return nodes

        math.extend(sequence())
        paragraph._p.append(math)
        paragraph.add_run(f"\t({self.equations})")

    def manuscript(self, path: Path) -> None:
        lines = path.read_text(encoding="utf-8").splitlines()
        index = 0
        while index < len(lines):
            line = lines[index].strip()
            if not line:
                index += 1
                continue
            marker = MARKER.fullmatch(line)
            if marker:
                if marker[1] == "FIGURE":
                    self.figure(marker[2])
                else:
                    self.pending_table = self.tables[marker[2]]
                index += 1
                continue
            if line.startswith("<!--"):
                index += 1
                continue
            if line.startswith("#"):
                level = len(line.split(" ")[0])
                self.heading(line[level + 1 :], level)
            elif line.startswith("|"):
                rows = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    raw = lines[index].strip()
                    if not re.fullmatch(r"[| :\-]+", raw):
                        rows.append([cell.strip() for cell in raw.strip("|").split("|")])
                    index += 1
                if not self.pending_table:
                    raise ValueError("Table needs named caption near " + self.last_heading)
                self.table(rows, self.pending_table)
                self.pending_table = None
                continue
            elif line.startswith("$$"):
                formula = line[2:]
                while not formula.endswith("$$"):
                    index += 1
                    formula += " " + lines[index].strip()
                self.equation(formula[:-2].strip())
            elif line.startswith("```"):
                index += 1
                while index < len(lines) and not lines[index].startswith("```"):
                    self.paragraph(lines[index], "CodeExample")
                    index += 1
            else:
                paragraph_lines = [line]
                while (
                    index + 1 < len(lines)
                    and lines[index + 1].strip()
                    and not re.match(r"[#|<]", lines[index + 1])
                ):
                    if lines[index + 1].startswith(("$$", "```", "- ")):
                        break
                    index += 1
                    paragraph_lines.append(lines[index].strip())
                text = " ".join(paragraph_lines)
                paragraph = self.paragraph(text)
                if len(text) < 140 and "[таблица:" in text:
                    paragraph.paragraph_format.keep_with_next = True
            index += 1

    def bibliography(self) -> None:
        paragraph = self.paragraph("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "Structural")
        self.bookmark(paragraph, "bibliography")
        for key, number in self.citations.items():
            self.paragraph(f"{number} {self.library[key]['bibliography']}", "Bibliography")

    def build(self, output: Path) -> dict[str, Any]:
        self.title_page()
        self.assignment()
        self.abstract()
        self.toc()
        for path in self.files:
            if path.parts[-2] != "appendices":
                self.manuscript(path)
        self.bibliography()
        for path in self.files:
            if path.parts[-2] == "appendices":
                self.manuscript(path)
        visible_text = "\n".join(
            "".join(paragraph.xpath(".//w:t/text() | .//m:t/text()"))
            for paragraph in self.document.element.xpath(".//w:p")
        )
        visible_words = len(re.findall(r"\b[\w-]+\b", visible_text))
        visible_characters = len(visible_text.replace("\n", " "))
        output.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(output)
        # Zip member timestamps are operational, so normalize them for repeatable builds.
        with ZipFile(output) as source:
            members = [(i.filename, source.read(i.filename)) for i in source.infolist()]
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
            for name, raw in members:
                if name == "docProps/app.xml":
                    app = ET.fromstring(raw)
                    namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
                    statistics = {
                        "Pages": self.pages,
                        "Words": visible_words,
                        "Characters": len(re.sub(r"\s", "", visible_text)),
                        "CharactersWithSpaces": visible_characters,
                    }
                    for tag, value in statistics.items():
                        node = app.find(namespace + tag)
                        if node is None:
                            node = ET.SubElement(app, namespace + tag)
                        node.text = str(value)
                    raw = ET.tostring(app, encoding="utf-8", xml_declaration=True)
                info = ZipInfo(name, (2026, 10, 16, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                target.writestr(info, raw)
        result = dict(
            docx=output.name,
            pages=self.pages,
            figures=len(self.figures),
            tables=len(self.tables),
            equations=self.equations,
            appendices=len([h for h in self.headings if h["title"].startswith("ПРИЛОЖЕНИЕ ")]),
            references=len(self.citations),
            words=visible_words,
            characters_with_spaces=visible_characters,
            characters_without_spaces=len(re.sub(r"\s", "", visible_text)),
            count_method="OOXML body visible text, tables and TOC included; Unicode word sequences",
            citation_order=self.citations,
            headings=self.headings,
            source_sha256={str(p.relative_to(ROOT)): sha(p) for p in self.files},
            source_words=len(re.findall(r"\b[\w-]+\b", self.text)),
            source_characters=len(self.text),
            docx_sha256=sha(output),
        )
        (THESIS / "internal/build-counts.json").write_text(
            json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        return result


def record_pagination(pdf: Path) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(pdf)
    texts = [page.extract_text() or "" for page in reader.pages]
    counts = read_json(THESIS / "internal/build-counts.json")
    headings: dict[str, int] = {}

    def normalize(value: str) -> str:
        return re.sub(r"\s+", "", value).replace("\u00ad", "")

    normalized = [normalize(text) for text in texts]
    for heading in counts["headings"] + [{"title": "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"}]:
        matches = [
            i + 1 for i, text in enumerate(normalized) if normalize(heading["title"]) in text
        ]
        if not matches:
            raise ValueError("Heading absent from render: " + heading["title"])
        headings[heading["title"]] = matches[-1]
    blank = [i + 1 for i, text in enumerate(texts) if not re.sub(r"[\d\s]", "", text)]
    if blank:
        raise ValueError("Blank rendered pages: " + str(blank))
    bibliography = headings["СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"]
    body_pages = bibliography - headings["ВВЕДЕНИЕ"]
    if body_pages < 90:
        raise ValueError("Main text below 90 pages")
    previous_path = THESIS / "internal/pagination.json"
    previous = read_json(previous_path) if previous_path.exists() else {}
    stable = previous.get("pages") == len(texts) and previous.get("headings") == headings
    result = dict(
        pages=len(texts),
        headings=headings,
        main_text_pages=body_pages,
        minimum_90="PASS",
        blank_pages=blank,
        provenance="Actual bundled LibreOffice PDF; last heading occurrence after TOC",
        pdf_sha256=sha(pdf),
        rendered_docx_sha256=counts["docx_sha256"],
        pagination_stable=stable,
    )
    previous_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return result


def render_for_review(renderer: Path, output: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    font = Path(
        manifest.get(
            "table_measurement_font", "/System/Library/Fonts/Supplemental/Times New Roman.ttf"
        )
    )
    if not font.is_file():
        raise FileNotFoundError("Times New Roman font required for layout review")
    temporary = THESIS / "tmp"
    temporary.mkdir(exist_ok=True)
    configuration = ET.Element("fontconfig")
    ET.SubElement(configuration, "dir").text = str(font.parent)
    ET.SubElement(configuration, "cachedir").text = str(temporary / "fontcache")
    fontconfig = temporary / "fontconfig.xml"
    ET.ElementTree(configuration).write(fontconfig, encoding="utf-8", xml_declaration=True)
    environment = os.environ.copy()
    environment["FONTCONFIG_FILE"] = str(fontconfig)
    environment["SAL_FONTPATH"] = str(font.parent)
    render = THESIS / "render"
    subprocess.run(
        [sys.executable, str(renderer), str(output), "--output_dir", str(render), "--emit_pdf"],
        env=environment,
        check=True,
    )
    result = record_pagination(render / (output.stem + ".pdf"))
    for page in render.glob("page-*.png"):
        if int(page.stem.split("-")[-1]) > result["pages"]:
            page.unlink()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, default=THESIS / "internal/build-manifest.json")
    parser.add_argument("--output", type=Path, default=THESIS / "ArchGuard_AI_VKR_MASTER.docx")
    parser.add_argument("--pages", type=int)
    parser.add_argument("--pagination-pdf", type=Path)
    parser.add_argument(
        "--render-with",
        type=Path,
        help="Bundled documents render_docx.py; never system LibreOffice",
    )
    args = parser.parse_args()
    if args.pagination_pdf:
        result = record_pagination(args.pagination_pdf)
        print(
            json.dumps(
                {
                    key: result[key]
                    for key in ("pages", "main_text_pages", "minimum_90", "pagination_stable")
                }
            )
        )
        return
    manifest = read_json(args.manifest)
    for attempt in range(4):
        result = ThesisBuilder(manifest, args.pages).build(args.output)
        if not args.render_with:
            break
        pagination = render_for_review(args.render_with, args.output, manifest)
        if pagination["pagination_stable"] and result["pages"] == pagination["pages"]:
            break
        if attempt == 3:
            raise RuntimeError("Pagination failed to stabilize in four renders")
    print(
        json.dumps(
            {
                k: result[k]
                for k in (
                    "pages",
                    "figures",
                    "tables",
                    "equations",
                    "appendices",
                    "references",
                    "source_words",
                )
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
